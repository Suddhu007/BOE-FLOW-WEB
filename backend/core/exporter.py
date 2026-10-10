import io

from openpyxl import Workbook


def _write_rows(sheet, records):
    columns = list(dict.fromkeys(key for record in records for key in record))
    if not columns:
        return
    sheet.append(columns)
    for record in records:
        sheet.append([record.get(column, "") for column in columns])


def build_excel(header, grouped_items, items):
    """Create the BOE export with openpyxl, avoiding the pandas runtime."""
    workbook = Workbook()
    parties = workbook.active
    parties.title = "Consignment & Parties"
    _write_rows(parties, [header])
    _write_rows(workbook.create_sheet("E-Way Bill Summary"), grouped_items)
    _write_rows(workbook.create_sheet("Line Items Audit"), items)

    total_assessable = sum(float(row.get("Assessable Value (CIF INR)", 0) or 0) for row in items)
    normalized = [{
        "BE_NO": header.get("BOE Number", ""),
        "BE_DATE": header.get("BOE Date", ""),
        "PORT_CODE": header.get("Port Code", ""),
        "GSTIN": header.get("Importer GSTIN", ""),
        "IMPORTER_NAME": header.get("Importer Name", ""),
        "SUPPLIER_NAME": header.get("Supplier Name", ""),
        "INVOICE_NUMBER": header.get("Invoice Number", ""),
        "TOTAL_DUTY": header.get("Total Duty", 0),
        "TOTAL_ASSESSABLE_VALUE": total_assessable,
        "ITEM_SN": item.get("Item No", ""),
        "ITEM_CTH": item.get("HSN Code", ""),
        "ITEM_DESCRIPTION": item.get("Description", ""),
        "ITEM_QUANTITY": item.get("Quantity", ""),
        "ITEM_UQC": item.get("UQC", ""),
        "ITEM_ASSESSABLE_VALUE": item.get("Assessable Value (CIF INR)", ""),
        "ITEM_TOTAL_DUTY": item.get("Item Total Duty", ""),
        "BCD_AMOUNT": item.get("BCD Amount", ""),
        "SWS_AMOUNT": item.get("SWS Amount", ""),
        "IGST_AMOUNT": item.get("IGST Amount", ""),
    } for item in items]
    _write_rows(workbook.create_sheet("SEPFUST Style Extract"), normalized)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
